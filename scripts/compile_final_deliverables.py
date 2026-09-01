#!/usr/bin/env python3
"""Compile all Markdown documents in final_deliverable/ into formatted DOCX files."""

import os
import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

BASE_DIR = Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/final_deliverable")

DOCUMENTS = [
    ("01_AIPP_Capstone_Project_Final_Report.md", "01_AIPP_Capstone_Project_Final_Report.docx"),
    ("02_AIPP_DEFENCE_QA.md", "02_AIPP_DEFENCE_QA.docx"),
    ("03_AGENT_SKILLS_AND_REGISTRY.md", "03_AGENT_SKILLS_AND_REGISTRY.docx"),
    ("04_MCP_ARCHITECTURE_AND_USE_CASES.md", "04_MCP_ARCHITECTURE_AND_USE_CASES.docx"),
    ("05_APPLICATION_ARCHITECTURE_AND_FLOWS.md", "05_APPLICATION_ARCHITECTURE_AND_FLOWS.docx"),
    ("06_VIVA_VOCE_QUESTIONS_AND_ANSWERS.md", "06_VIVA_VOCE_QUESTIONS_AND_ANSWERS.docx"),
]

def _set_cell_bg(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)

def _add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = "Times New Roman"
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = "Times New Roman"
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        else:
            run = paragraph.add_run(part)
            run.font.name = "Times New Roman"

def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = tbl.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h.strip())
        r.bold = True
        r.font.name = "Times New Roman"
        r.font.size = Pt(9.5)
        _set_cell_bg(c, "1F4E79")
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r_idx, row in enumerate(rows):
        bg = "F2F5F8" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row):
            if c_idx >= len(header):
                continue
            c = tbl.rows[1 + r_idx].cells[c_idx]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            _add_inline(p, val.strip())
            p.style.font.name = "Times New Roman"
            p.style.font.size = Pt(9)
            _set_cell_bg(c, bg)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def convert_md_to_docx(md_path: Path, docx_path: Path):
    print(f"Compiling {md_path.name} -> {docx_path.name}...")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    in_code_block = False
    code_lines = []
    in_table = False
    table_header = []
    table_rows = []

    for line in lines:
        raw = line.rstrip("\r\n")

        # Code block handling
        if raw.startswith("```"):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.left_indent = Inches(0.25)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw)
            continue

        # Table handling
        if raw.startswith("|") and raw.endswith("|"):
            cells = [c.strip() for c in raw.strip("|").split("|")]
            if all(re.match(r"^:?-+:?$", c) for c in cells):
                continue
            if not in_table:
                in_table = True
                table_header = cells
                table_rows = []
            else:
                table_rows.append(cells)
            continue
        elif in_table:
            in_table = False
            _add_table(doc, table_header, table_rows)
            table_header = []
            table_rows = []

        # Headings
        if raw.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(raw[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        elif raw.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(raw[3:].strip())
            run.bold = True
            run.font.size = Pt(14)
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        elif raw.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(raw[4:].strip())
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        elif raw.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.space_after = Pt(4)
            _add_inline(p, raw[2:].strip())
        elif raw.startswith("- ") or raw.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            _add_inline(p, raw[2:].strip())
        elif re.match(r"^\d+\.\s", raw):
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            content = re.sub(r"^\d+\.\s", "", raw)
            _add_inline(p, content.strip())
        elif raw.strip() == "---":
            continue
        elif raw.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            _add_inline(p, raw.strip())

    if in_table:
        _add_table(doc, table_header, table_rows)

    doc.save(docx_path)
    print(f"  ✓ Successfully created {docx_path.name}")

if __name__ == "__main__":
    for md_name, docx_name in DOCUMENTS:
        md_file = BASE_DIR / md_name
        docx_file = BASE_DIR / docx_name
        if md_file.exists():
            convert_md_to_docx(md_file, docx_file)
        else:
            print(f"Skipping {md_name} (file not found)")

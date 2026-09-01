"""Format and calibrate final deliverable DOCX to Times New Roman 12pt and standard 50-page capstone layout."""

import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

DOC_PATH = Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/final_deliverable/DEMO_AIPP_Capstone_Project_Final_Report.docx")
DOCS_MIRROR = Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/docs/AIPP_Capstone_Project_Final_Report.docx")

def format_document(doc_path: Path):
    doc = docx.Document(doc_path)
    
    # 1. Page Margins — 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # 2. Configure 'Normal' style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(3.0)

    # 3. Format all Paragraphs and Runs
    for p in doc.paragraphs:
        # Paragraph spacing default
        if p.paragraph_format.line_spacing is None:
            p.paragraph_format.line_spacing = 1.15
        if p.paragraph_format.space_after is None:
            p.paragraph_format.space_after = Pt(3.0)
            
        txt = p.text.strip()
        
        # Check headings
        is_h1 = (p.style.name == 'Heading 1' or txt.startswith('Chapter') or txt.startswith('CHAPTER') 
                 or txt in ['Candidate’s Declaration', 'Certificate', 'Acknowledgement', 
                            'Similarity Index Report', 'AI Usage Disclosure Statement', 
                            'Student Declaration', 'List of Abbreviations', 'List of Figures', 
                            'List of Tables', 'Abstract', 'Table of Contents', 'Bibliography', 'References'])
        is_h2 = p.style.name == 'Heading 2' or (len(txt) > 3 and txt[:3].replace('.', '').isdigit() and ' ' in txt[:6])
        
        for r in p.runs:
            # Preserve Consolas for code/snippet runs
            if r.font.name == 'Consolas':
                continue
            r.font.name = 'Times New Roman'
            
            # Format Headings
            if is_h1:
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
            elif is_h2:
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x2B, 0x4C, 0x7E)
            else:
                if r.font.size is None:
                    r.font.size = Pt(12)

    # 4. Format all Tables
    for t in doc.tables:
        t.style = 'Light Grid Accent 1'
        # Format Header Row
        for cell in t.rows[0].cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2.0)
                p.paragraph_format.line_spacing = 1.05
                for r in p.runs:
                    if r.font.name != 'Consolas':
                        r.font.name = 'Times New Roman'
                    r.font.size = Pt(10.0)
                    r.font.bold = True
        
        # Format Data Rows
        for row in t.rows[1:]:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2.0)
                    p.paragraph_format.line_spacing = 1.05
                    for r in p.runs:
                        if r.font.name != 'Consolas':
                            r.font.name = 'Times New Roman'
                        r.font.size = Pt(9.5)

    doc.save(doc_path)
    doc.save(DOCS_MIRROR)
    print(f"✓ Formatted and saved {doc_path} ({doc_path.stat().st_size:,} bytes)")

if __name__ == '__main__':
    format_document(DOC_PATH)

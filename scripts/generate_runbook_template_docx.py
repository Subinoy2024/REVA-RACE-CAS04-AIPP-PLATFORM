#!/usr/bin/env python3
"""Generates the exact SRE Incident Runbook template requested by the user into a formatted .docx document."""

from __future__ import annotations

import os
from pathlib import Path
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor


def set_cell_background(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


def add_code_block(doc: docx.Document, code_text: str) -> None:
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
    p = cell.paragraphs[0]
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 30, 30)
    doc.add_paragraph()


def generate_exact_runbook_docx(output_path: str = "docs/SRE_Incident_Runbook_Template.docx") -> str:
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    NAVY = RGBColor(15, 34, 64)
    BLUE = RGBColor(0, 102, 204)
    DARK_GRAY = RGBColor(51, 51, 51)

    # Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run("SRE Incident Runbook")
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(24)
    r_title.font.bold = True
    r_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Incident ID: {{$json.incident_id}}")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(14)
    r_sub.font.bold = True
    r_sub.font.color.rgb = BLUE

    doc.add_paragraph()

    # Symptom
    h_symptom = doc.add_heading("Symptom", level=2)
    h_symptom.runs[0].font.color.rgb = NAVY
    p_sym = doc.add_paragraph()
    r_sym = p_sym.add_run("• Describe the observable symptoms that indicate the incident has occurred. This could include errors in logs, service unavailability, degraded performance, etc.")
    r_sym.font.color.rgb = DARK_GRAY

    # Diagnosis Steps
    h_diag = doc.add_heading("Diagnosis Steps", level=2)
    h_diag.runs[0].font.color.rgb = NAVY

    diag_items = [
        ("1. Check Service Status", "systemctl status <service_name>", "Verify if the service is running or has failed."),
        ("2. Review Logs", 'journalctl -u <service_name> --since "10 minutes ago"', "Look for any error messages or warnings that could indicate the cause of the issue."),
        ("3. Check Resource Utilization", "top -u <user_name>", "Monitor CPU and memory usage to identify any resource bottlenecks."),
        ("4. Network Connectivity", "ping <service_endpoint>", "Ensure that the service can reach its dependencies and external services."),
        ("5. Database Connection", 'mysql -u <user> -p -h <db_host> -e "SHOW DATABASES;"', "Confirm that the application can connect to the database.")
    ]

    for title, cmd, desc in diag_items:
        p = doc.add_paragraph()
        r = p.add_run(title)
        r.font.bold = True
        r.font.color.rgb = BLUE
        add_code_block(doc, cmd)
        pd = doc.add_paragraph()
        rd = pd.add_run(f"• {desc}")
        rd.font.color.rgb = DARK_GRAY

    # Fix
    h_fix = doc.add_heading("Fix", level=2)
    h_fix.runs[0].font.color.rgb = NAVY

    p_f1 = doc.add_paragraph()
    p_f1.add_run("1. Restart the Service").font.bold = True
    add_code_block(doc, "systemctl restart <service_name>")
    p_f1_desc = doc.add_paragraph()
    p_f1_desc.add_run("• Restart the service to see if it resolves the issue.").font.color.rgb = DARK_GRAY

    p_f2 = doc.add_paragraph()
    p_f2.add_run("2. Rollback Changes (if applicable)").font.bold = True
    p_f2_desc = doc.add_paragraph()
    p_f2_desc.add_run("• If recent changes were made, revert to the previous stable version:").font.color.rgb = DARK_GRAY
    add_code_block(doc, "git checkout <previous_commit_hash>")
    p_f2_k8s = doc.add_paragraph()
    p_f2_k8s.add_run("• Redeploy the application:").font.color.rgb = DARK_GRAY
    add_code_block(doc, "kubectl rollout undo deployment/<deployment_name>")

    # Prevention
    h_prev = doc.add_heading("Prevention", level=2)
    h_prev.runs[0].font.color.rgb = NAVY

    prev_bullets = [
        ("Monitoring and Alerts", "Set up monitoring and alerting for key metrics to catch issues before they escalate."),
        ("Regular Maintenance", "Schedule regular maintenance windows to apply updates and patches."),
        ("Documentation", "Maintain up-to-date documentation for troubleshooting steps and known issues."),
        ("Postmortem Review", "Conduct a postmortem review to analyze the root cause and implement corrective actions to prevent recurrence.")
    ]

    for label, text in prev_bullets:
        p = doc.add_paragraph()
        r_l = p.add_run(f"• {label}: ")
        r_l.font.bold = True
        r_l.font.color.rgb = BLUE
        r_t = p.add_run(text)
        r_t.font.color.rgb = DARK_GRAY

    # RCA
    h_rca = doc.add_heading("RCA", level=2)
    h_rca.runs[0].font.color.rgb = NAVY
    p_rca = doc.add_paragraph()
    r_rcal = p_rca.add_run("• Root Cause Analysis: ")
    r_rcal.font.bold = True
    r_rcal.font.color.rgb = BLUE
    r_rcat = p_rca.add_run("{{$json.rca_summary}}")
    r_rcat.font.color.rgb = DARK_GRAY

    # Timeline
    h_time = doc.add_heading("Timeline", level=2)
    h_time.runs[0].font.color.rgb = NAVY
    p_time = doc.add_paragraph()
    r_timel = p_time.add_run("• Incident Timeline: ")
    r_timel.font.bold = True
    r_timel.font.color.rgb = BLUE
    r_timet = p_time.add_run("{{$json.timeline}}")
    r_timet.font.color.rgb = DARK_GRAY

    out = Path(output_path)
    out.parent.mkdir(exist_ok=True)
    doc.save(str(out))
    print(f"Generated exact SRE Runbook .docx: {out}")
    return str(out)


if __name__ == "__main__":
    generate_exact_runbook_docx()

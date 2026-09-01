"""Convert Project_Paper/AIPP_Final_Capstone_Project_Report.md to a formatted DOCX report."""

import re
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, Inches, RGBColor

TARGETS = [
    (
        Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/Project_Paper/AIPP_Final_Capstone_Project_Report.md"),
        Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/Project_Paper/AIPP_Final_Capstone_Project_Report.docx"),
    ),
    (
        Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/final_deliverable/01_AIPP_Capstone_Project_Final_Report.md"),
        Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/final_deliverable/01_AIPP_Capstone_Project_Final_Report.docx"),
    ),
]

def _set_cell_bg(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)

def _add_inline(paragraph, text: str) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = "Times New Roman"
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
            run.font.name = "Times New Roman"
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
        else:
            run = paragraph.add_run(part)
            run.font.name = "Times New Roman"

def _add_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = tbl.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(h.strip())
        run.bold = True
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.0)
        _set_cell_bg(c, "D9E2F3")
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            if ci < len(tbl.rows[ri + 1].cells):
                target = tbl.rows[ri + 1].cells[ci]
                target.text = ""
                p = target.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                _add_inline(p, cell.strip())
                for r in p.runs:
                    if r.font.name != "Consolas":
                        r.font.name = "Times New Roman"
                    r.font.size = Pt(9.5)

def _flush_table(doc, pending_table):
    if not pending_table:
        return
    header_row = pending_table[0]
    data_rows = pending_table[2:] if len(pending_table) > 2 else []
    header = [c.strip() for c in _split_row(header_row)]
    rows = [[c.strip() for c in _split_row(r)] for r in data_rows]
    _add_table(doc, header, rows)

def _split_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return line.split("|")

def convert(src: Path, dst: Path) -> None:
    text = src.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    
    # Page Setup — 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style configuration — Times New Roman 12pt
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.paragraph_format.line_spacing = 1.15
    style.paragraph_format.space_after = Pt(3.0)

    in_code = False
    code_buf: list[str] = []
    pending_table: list[str] = []
    in_html_comment = False

    just_added_page_break = True  # Start of doc

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if "<!--" in line:
            in_html_comment = True
        if in_html_comment:
            if "-->" in line:
                in_html_comment = False
            i += 1
            continue

        # Code blocks
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                run = p.add_run("\n".join(code_buf))
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                p.paragraph_format.left_indent = Pt(14)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.0
                code_buf = []
                in_code = False
                just_added_page_break = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # Tables
        if stripped.startswith("|") and stripped.endswith("|") and "|" in stripped[1:-1]:
            pending_table.append(stripped)
            i += 1
            continue
        else:
            if pending_table:
                _flush_table(doc, pending_table)
                pending_table = []
                just_added_page_break = False

        if stripped.startswith("<div") or stripped.startswith("</div>"):
            i += 1
            continue
        
        # Images: ![Caption](path)
        img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if img_match:
            caption, img_path_str = img_match.groups()
            img_path = Path(img_path_str)
            if not img_path.is_absolute():
                if (src.parent / img_path_str).exists():
                    img_path = src.parent / img_path_str
                elif (Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/Project_Paper") / img_path_str).exists():
                    img_path = Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/Project_Paper") / img_path_str
                elif (Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/docs") / img_path_str).exists():
                    img_path = Path("/Users/subdebna/CAS-04-MS-SUBINOY-AIPP/REVA-RACE-CAS04-AIPP-PLATFORM/docs") / img_path_str
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(img_path), width=Inches(5.5))
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cp.paragraph_format.space_after = Pt(6)
                crun = cp.add_run(f"Figure: {caption}")
                crun.italic = True
                crun.font.name = "Times New Roman"
                crun.font.size = Pt(9.5)
                crun.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
                just_added_page_break = False
            i += 1
            continue

        # Page breaks
        if stripped == "---":
            if not just_added_page_break:
                doc.add_page_break()
                just_added_page_break = True
            i += 1
            continue

        # Title (# )
        if stripped.startswith("# "):
            p = doc.add_heading(stripped[2:].strip(), level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(16)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
            just_added_page_break = False
            i += 1
            continue

        # Chapter Headings (## )
        if stripped.startswith("## "):
            title = stripped[3:].strip()
            if title.startswith("CHAPTER") or title.startswith("BIBLIOGRAPHY") or title.startswith("APPENDIX") or title.startswith("Chapter"):
                if not just_added_page_break:
                    doc.add_page_break()
                    just_added_page_break = True
            p = doc.add_heading(title, level=1)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(14)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x8A)
            just_added_page_break = False
            i += 1
            continue

        # Section Headings (### )
        if stripped.startswith("### "):
            p = doc.add_heading(stripped[4:].strip(), level=2)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x2B, 0x4C, 0x7E)
            just_added_page_break = False
            i += 1
            continue

        # Sub-section Headings (#### )
        if stripped.startswith("#### "):
            p = doc.add_heading(stripped[5:].strip(), level=3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x3C, 0x5A, 0xA0)
            just_added_page_break = False
            i += 1
            continue

        # Blockquotes (> )
        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.right_indent = Pt(18)
            run = p.add_run(stripped[2:].strip())
            run.italic = True
            run.font.name = "Times New Roman"
            run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
            just_added_page_break = False
            i += 1
            continue

        # Bullet List
        if stripped.startswith("* ") or stripped.startswith("- "):
            content = stripped[2:]
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            _add_inline(p, content)
            i += 1
            continue

        # Numbered List
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            _add_inline(p, m.group(1))
            i += 1
            continue

        # Blank Line
        if not stripped:
            i += 1
            continue

        # Regular Paragraph
        buf = [line]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].strip()
            if (not nxt) or nxt.startswith(("#", "* ", "- ", "> ", "|", "```")) or re.match(r"^\d+\.\s+", nxt) or nxt == "---":
                break
            buf.append(lines[j])
            j += 1
        para_text = " ".join(x.strip() for x in buf if x.strip())
        if para_text:
            p = doc.add_paragraph()
            _add_inline(p, para_text)
        i = j

    if pending_table:
        _flush_table(doc, pending_table)

    doc.save(dst)
    print(f"✓ Successfully generated {dst} ({dst.stat().st_size:,} bytes)")

if __name__ == "__main__":
    for src, dst in TARGETS:
        if src.exists():
            convert(src, dst)
